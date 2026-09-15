from __future__ import annotations

import csv
from pathlib import Path

from docling.document_converter import DocumentConverter


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".csv",
    ".txt",
    ".md",
}


def _parse_with_docling(input_file: Path):
    """Convert PDF/DOCX with Docling and return the Docling document."""
    converter = DocumentConverter()
    result = converter.convert(str(input_file))
    return result.document


def _docling_pages(document) -> list[dict]:
    """Extract text grouped by source page from a Docling document."""
    page_text: dict[int, list[str]] = {}

    for item, _ in document.iterate_items():
        if not item.prov:
            continue

        page_number = item.prov[0].page_no
        text = getattr(item, "text", None)

        if text:
            page_text.setdefault(page_number, []).append(text)

    pages = []
    for page_number in sorted(page_text):
        text = "\n".join(page_text[page_number]).strip()
        if text:
            pages.append({"page": page_number, "text": text})

    return pages


def _spreadsheet_pages(input_file: Path) -> list[dict]:
    """Read XLSX/XLS and represent each worksheet as one logical page."""
    suffix = input_file.suffix.lower()

    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(
            input_file,
            read_only=True,
            data_only=True,
        )

        pages = []
        for page_number, sheet in enumerate(workbook.worksheets, start=1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [
                    "" if value is None else str(value).strip()
                    for value in row
                ]
                if any(values):
                    rows.append(values)

            if not rows:
                continue

            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]

            # Keep the spreadsheet faithful and easy to inspect as Markdown.
            header = rows[0]
            separator = ["---"] * width
            markdown_rows = [
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(separator) + " |",
            ]
            markdown_rows.extend(
                "| " + " | ".join(row) + " |"
                for row in rows[1:]
            )

            text = f"## Sheet: {sheet.title}\n\n" + "\n".join(markdown_rows)
            pages.append(
                {
                    "page": page_number,
                    "section": sheet.title,
                    "text": text,
                }
            )

        workbook.close()
        return pages

    # .xls is handled separately because openpyxl does not support it.
    import xlrd

    workbook = xlrd.open_workbook(input_file, on_demand=True)
    pages = []

    for page_number, sheet_name in enumerate(
        workbook.sheet_names(),
        start=1,
    ):
        sheet = workbook.sheet_by_name(sheet_name)
        rows = []

        for row_index in range(sheet.nrows):
            values = [
                "" if sheet.cell_value(row_index, col_index) is None
                else str(sheet.cell_value(row_index, col_index)).strip()
                for col_index in range(sheet.ncols)
            ]
            if any(values):
                rows.append(values)

        if not rows:
            continue

        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]

        header = rows[0]
        separator = ["---"] * width
        markdown_rows = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(separator) + " |",
        ]
        markdown_rows.extend(
            "| " + " | ".join(row) + " |"
            for row in rows[1:]
        )

        text = f"## Sheet: {sheet_name}\n\n" + "\n".join(markdown_rows)
        pages.append(
            {
                "page": page_number,
                "section": sheet_name,
                "text": text,
            }
        )

    workbook.release_resources()
    return pages


def _csv_pages(input_file: Path) -> list[dict]:
    """Represent a CSV as one logical page while preserving cell values."""
    with input_file.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))

    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return []

    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]

    header = rows[0]
    separator = ["---"] * width
    markdown_rows = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    markdown_rows.extend(
        "| " + " | ".join(row) + " |"
        for row in rows[1:]
    )

    return [{"page": 1, "text": "\n".join(markdown_rows)}]


def _plain_text_pages(input_file: Path) -> list[dict]:
    text = input_file.read_text(encoding="utf-8").strip()
    return [{"page": 1, "text": text}] if text else []


def parse_document_pages(input_file: Path) -> list[dict]:
    """
    Parse a supported document and return logical pages/sections.

    PDF/DOCX:
        Docling extraction with source-page provenance where available.

    XLSX/XLS:
        One logical page per worksheet.

    CSV/TXT/MD:
        One logical page for the file.
    """
    input_file = Path(input_file)
    suffix = input_file.suffix.lower()

    if suffix in {".pdf", ".docx"}:
        document = _parse_with_docling(input_file)
        pages = _docling_pages(document)

        # Some DOCX files do not expose physical page provenance.
        if not pages:
            markdown = document.export_to_markdown().strip()
            if markdown:
                pages = [{"page": 1, "text": markdown}]

        return pages

    if suffix in {".xlsx", ".xls"}:
        return _spreadsheet_pages(input_file)

    if suffix == ".csv":
        return _csv_pages(input_file)

    if suffix in {".txt", ".md"}:
        return _plain_text_pages(input_file)

    raise ValueError(
        f"Unsupported document type: {input_file.suffix}"
    )


def parse_document(
    input_file: Path,
    output_dir: Path,
) -> Path:
    """Parse one supported file and save its faithful Markdown form."""
    input_file = Path(input_file)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"Parsing: {input_file.name}")

    pages = parse_document_pages(input_file)

    markdown_parts = []
    for page in pages:
        if page.get("section"):
            markdown_parts.append(f"<!-- Section: {page['section']} -->")
        markdown_parts.append(page["text"])

    markdown = "\n\n".join(markdown_parts).strip()

    output_file = output_dir / f"{input_file.stem}.md"
    output_file.write_text(markdown + "\n", encoding="utf-8")

    print(
        f"Saved: {output_file.name} "
        f"({len(markdown):,} characters, {len(pages)} logical page(s))"
    )

    return output_file
